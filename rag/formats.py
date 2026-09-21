"""Bounded extraction of structured text and modern Office documents."""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile

from .models import RagError

OFFICE_BYTES = 20 * 1024 * 1024


def decode(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise RagError("Text and structured data files must use UTF-8 encoding.") from None
    if "\x00" in text:
        raise RagError("Binary content is not supported as a text file.")
    return text


def flatten(value, path="", depth=0):
    if depth > 32:
        raise RagError("JSON nesting exceeds 32 levels. Simplify the document.")
    if isinstance(value, dict):
        for key, item in value.items():
            yield from flatten(item, f"{path}.{key}" if path else key, depth + 1)
    elif isinstance(value, list):
        for number, item in enumerate(value):
            yield from flatten(item, f"{path}[{number}]", depth + 1)
    elif value is not None:
        yield (f"{path}: " if path else "") + (json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value)


def parse_json(text, where):
    def pairs(items):
        result = {}
        for key, value in items:
            if key in result:
                raise ValueError
            result[key] = value
        return result
    def reject_constant(value):
        raise ValueError
    try:
        return json.loads(text, object_pairs_hook=pairs, parse_constant=reject_constant)
    except (ValueError, RecursionError):
        raise RagError(f"Invalid JSON in {where}. Check syntax, duplicate keys, and nesting.") from None


def json_record(record, where):
    if isinstance(record, dict) and "text" in record:
        if not isinstance(record["text"], str) or not record["text"].strip():
            raise RagError(f"The text field in {where} must be a non-empty string.")
        location = where
        for key in ("source_doc", "chunk_id"):
            if key in record:
                if not isinstance(record[key], (str, int)):
                    raise RagError(f"The {key} field in {where} must be text or a number.")
                location += f" · {key}: {str(record[key])[:200]}"
        # Labels, including 'attacked', are metadata, never a trust decision.
        extra = {key: value for key, value in record.items() if key not in {"text", "source_doc", "chunk_id", "attacked"}}
        return "\n".join([record["text"], *flatten(extra)]), location
    text = "\n".join(flatten(record))
    if not text.strip():
        raise RagError(f"No readable content in {where}.")
    return text, where


def table_rows(rows, location):
    iterator = iter(rows)
    header = next(iterator, None)
    if header is None:
        return
    if len(header) > 1000:
        raise RagError("The table exceeds 1,000 columns. Split it into smaller files.")
    header = [str(v) if v is not None and str(v).strip() else f"Column {i + 1}" for i, v in enumerate(header)]
    for number, row in enumerate(iterator, 2):
        if number > 100000 or len(row) > 1000:
            raise RagError("The table is too complex. Use at most 100,000 rows and 1,000 columns.")
        if not any(v is not None and str(v).strip() for v in row):
            continue
        text = "\n".join(f"{header[i] if i < len(header) else 'Column ' + str(i + 1)}: {value}"
                         for i, value in enumerate(row) if value is not None)
        yield None, text, f"{location} · row {number}"


def check_office(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 5000 or sum(m.file_size for m in members) > OFFICE_BYTES:
                raise RagError("Office content exceeds the 20 MB decompressed limit or is too complex.")
            if any(m.flag_bits & 1 for m in members):
                raise RagError("Encrypted Office files are not supported.")
            if any("vbaproject" in m.filename.lower() for m in members):
                raise RagError("Macro-enabled Office files are not supported.")
            # Read with a bounded cumulative budget before handing off to a parser.
            total = 0
            for member in members:
                with archive.open(member) as stream:
                    while True:
                        data = stream.read(min(65536, OFFICE_BYTES - total + 1))
                        if not data:
                            break
                        total += len(data)
                        if total > OFFICE_BYTES:
                            raise RagError("Office content exceeds the 20 MB decompressed limit.")
                if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                    # Reject huge sparse coordinates before openpyxl allocates a row.
                    from defusedxml.ElementTree import iterparse
                    with archive.open(member) as stream:
                        for _, element in iterparse(stream, events=("end",)):
                            if element.tag.endswith("}row") and int(element.get("r", "1")) > 100000:
                                raise RagError("The workbook exceeds 100,000 rows. Split it into smaller files.")
                            if element.tag.endswith("}c"):
                                match = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", element.get("r", "A1"))
                                if not match:
                                    raise RagError("The workbook contains invalid cell references.")
                                column = 0
                                for character in match[1]:
                                    column = column * 26 + ord(character) - ord("A") + 1
                                if column > 1000 or int(match[2]) > 100000:
                                    raise RagError("The workbook exceeds 100,000 rows or 1,000 columns. Split it into smaller files.")
                            element.clear()
    except RagError:
        raise
    except Exception:
        raise RagError("The Office file is corrupt, encrypted, or unsupported.") from None


def extract(suffix, content):
    """Yield (page, text, source location); never execute uploaded content."""
    if suffix in {"json", "jsonl"}:
        text = decode(content)
        if suffix == "jsonl":
            for line, value in enumerate(text.splitlines(), 1):
                if value.strip():
                    record_text, location = json_record(parse_json(value, f"line {line}"), f"line {line}")
                    yield None, record_text, location
        else:
            records = parse_json(text, "document")
            if not isinstance(records, (dict, list)):
                raise RagError("JSON must contain an object or an array of records.")
            records = records if isinstance(records, list) else [records]
            for number, record in enumerate(records, 1):
                record_text, location = json_record(record, f"record {number}")
                yield None, record_text, location
    elif suffix in {"csv", "tsv"}:
        try:
            yield from table_rows(csv.reader(io.StringIO(decode(content)), delimiter="," if suffix == "csv" else "\t", strict=True), "table")
        except csv.Error:
            raise RagError("The delimited table is malformed or contains an oversized field.") from None
    else:
        check_office(content)
        try:
            if suffix == "docx":
                from docx import Document
                from docx.table import Table
                document = Document(io.BytesIO(content))
                for number, block in enumerate(document.iter_inner_content(), 1):
                    if isinstance(block, Table):
                        # Word tables do not necessarily contain a header row.
                        for row_number, row in enumerate(block.rows, 1):
                            yield None, " | ".join(cell.text for cell in row.cells), f"table at block {number} · row {row_number}"
                    elif block.text.strip():
                        yield None, block.text, f"paragraph {number}"
            elif suffix == "xlsx":
                from openpyxl import load_workbook
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
                try:
                    if len(workbook.worksheets) > 100:
                        raise RagError("The workbook exceeds 100 sheets. Split it into smaller files.")
                    for sheet in workbook.worksheets:
                        # Ignore untrusted dimension hints; iterate the actual XML rows.
                        sheet.reset_dimensions()
                        yield from table_rows(sheet.iter_rows(values_only=True), f"sheet {sheet.title}")
                finally:
                    workbook.close()
        except RagError:
            raise
        except Exception:
            raise RagError("The Office file could not be read. Use a valid unencrypted DOCX or XLSX file.") from None
