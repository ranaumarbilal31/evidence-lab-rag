import io
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from rag.ingest import ingest, BYOK_LIMITS
from rag.models import Evidence, RagError
from rag.pipeline import check_sources


@pytest.mark.parametrize("name,data,expected,location", [
    ("policy.json", b'{"policy":{"leave":20},"groups":["staff"]}', "policy.leave: 20", "record 1"),
    ("policy.jsonl", b'{"text":"Leave is 20 days"}\n{"text":"Refund is 30 days"}', "Leave is 20 days", "line 1"),
    ("policy.csv", b'group,days\nstaff,20\n', "days: 20", "row 2"),
    ("policy.tsv", b'group\tdays\nstaff\t20\n', "days: 20", "row 2"),
    ("policy.md", b'# Leave\n20 days', "20 days", None),
    ("policy.txt", b'Leave is 20 days', "20 days", None),
])
def test_formats_preserve_text_and_provenance(name, data, expected, location):
    chunks, _ = ingest([(name, data)])
    assert expected in chunks[0].text
    assert location is None or location in chunks[0].location
    sources = check_sources([Evidence(chunk_id=chunks[0].id, quote=expected)], {c.id: c for c in chunks})
    assert sources[0].get("location") == chunks[0].location


def test_json_existing_schema_duplicate_external_ids_and_labels():
    data = json.dumps([{"chunk_id": "same", "text": "First policy", "source_doc": "A", "attacked": False},
                       {"chunk_id": "same", "text": "Second policy", "source_doc": "B", "attacked": True}]).encode()
    chunks, _ = ingest([("records.json", data)])
    assert len({c.id for c in chunks}) == 2
    assert "source_doc: A" in chunks[0].location
    assert "attacked" not in chunks[0].text and "attacked" not in chunks[1].text


def test_existing_corpus_is_accepted_with_personal_limits():
    path = Path(__file__).parents[1] / "rag-safety-wall" / "corpus.json"
    chunks, _ = ingest([("corpus.json", path.read_bytes())], BYOK_LIMITS)
    assert len(chunks) >= 1034 and len({c.id for c in chunks}) == len(chunks)
    with pytest.raises(RagError, match="50 chunks"):
        ingest([("corpus.json", path.read_bytes())])


@pytest.mark.parametrize("name,data,match", [
    ("bad.json", b'{"a":1,"a":2}', "Invalid JSON"),
    ("bad.json", b'{"text":42}', "text field"),
    ("bad.json", b'"scalar"', "object or an array"),
    ("bad.json", b'{"n":NaN}', "Invalid JSON"),
    ("bad.jsonl", b'{"text":"ok"}\nnot-json', "line 2"),
    ("bad.json", b'[{}]', "No readable content"),
    ("bad.json", b'\xff', "UTF-8"),
    ("bad.csv", b'header\n"unfinished', "malformed"),
    ("bad.docx", b'not a zip', "corrupt"),
    ("bad.xlsx", b'not a zip', "corrupt"),
    ("bad.doc", b'legacy', "not supported"),
    ("bad.zip", b'archive', "not supported"),
])
def test_invalid_data_has_readable_error(name, data, match):
    with pytest.raises(RagError, match=match): ingest([(name, data)])


def test_deep_json_is_rejected_without_crash():
    data = ('{"nested":' * 40 + '"value"' + '}' * 40).encode()
    with pytest.raises(RagError, match="nesting"):
        ingest([("deep.json", data)])


def test_docx_paragraphs_and_tables():
    doc = Document()
    doc.add_paragraph("Employees receive 20 days annual leave.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Policy"
    table.cell(0, 1).text = "Days"
    table.cell(1, 0).text = "Refund"
    table.cell(1, 1).text = "30"
    data = io.BytesIO(); doc.save(data)
    chunks, _ = ingest([("policy.docx", data.getvalue())])
    assert "20 days" in chunks[0].text and "paragraph" in chunks[0].location
    assert any("Refund" in c.text and "row 2" in c.location for c in chunks)


def test_xlsx_values_and_sheet_rows_no_formula_execution():
    book = Workbook(); sheet = book.active; sheet.title = "Policies"
    sheet.append(["Policy", "Days", "Formula"])
    sheet.append(["Refund", 30, "=1+1"])
    data = io.BytesIO(); book.save(data)
    chunks, _ = ingest([("policy.xlsx", data.getvalue())])
    assert "Days: 30" in chunks[0].text and "sheet Policies" in chunks[0].location
    assert "row 2" in chunks[0].location and "=1+1" not in chunks[0].text


def test_office_decompressed_limit():
    data = io.BytesIO()
    with zipfile.ZipFile(data, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (20 * 1024 * 1024 + 1))
    with pytest.raises(RagError, match="20 MB"):
        ingest([("bomb.docx", data.getvalue())])


def test_oversized_upload_rejected_in_personal_mode():
    with pytest.raises(RagError, match="upload size"):
        ingest([("large.txt", b"x" * (2 * 1024 * 1024 + 1))], BYOK_LIMITS)
